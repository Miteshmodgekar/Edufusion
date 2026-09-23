"""Marks Blueprint — VTU CIE marks management.

Supports 5 subject types matching the college Excel format:
  ipcc_theory  : Q1-Q4 (a,b,c,d) x2 IAs, scale to 25M, Assign 25M
  cc_theory    : Q1-Q4 (a,b,c,d) x2 IAs, scale to 15M, Assign 20M
  ipcc_lab     : Q1-Q4 x2 IAs (15M) + Lab IA 25M + Assign 10M
  cc_activity  : Module 1-5 (20M each)
  cc_oe        : Gen 30M + CIE 20M

Two entry modes:
  1. Excel upload (download template → fill → upload)
  2. Direct in-app entry (like attendance — no Excel needed)
"""
from flask import Blueprint, request, jsonify, render_template, send_file
from flask_login import login_required, current_user
from extensions import db
from models.internal_marks import InternalMarks, _q_total, ASSIGN_MAX
from models.user import User
import io
from datetime import datetime
from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

marks_bp = Blueprint("marks", __name__)

# ── Helpers ───────────────────────────────────────────────────────────────────

def _fac_ok():
    return current_user.is_faculty or current_user.is_hod or current_user.is_admin

def _f(v):
    try: return round(float(v), 2) if v not in (None, "", "nan") else None
    except: return None

TYPE_LABELS = {
    'ipcc_theory': 'IPCC — Theory',
    'cc_theory':   'CC — Theory',
    'ipcc_lab':    'IPCC — Lab (Theory + Lab IA)',
    'cc_activity': 'CC — Activity (5 Modules)',
    'cc_oe':       'CC — Open Elective',
}

SCALE_MAX = {'ipcc_theory': 25, 'cc_theory': 15, 'ipcc_lab': 15}

# ── Style helpers ─────────────────────────────────────────────────────────────

def _mk_styles():
    HDR   = PatternFill("solid", fgColor="1F3864")
    Q1F   = PatternFill("solid", fgColor="D9EAD3")
    Q2F   = PatternFill("solid", fgColor="CFE2F3")
    Q3F   = PatternFill("solid", fgColor="FCE5CD")
    Q4F   = PatternFill("solid", fgColor="EAD1DC")
    TOTF  = PatternFill("solid", fgColor="FFF2CC")
    LABF  = PatternFill("solid", fgColor="E8D5FF")
    MODF  = [PatternFill("solid", fgColor=c) for c in
             ["D9EAD3","CFE2F3","FCE5CD","EAD1DC","FFE4B5"]]
    FW    = Font(bold=True, color="FFFFFF", size=9)
    FB    = Font(bold=True, color="000000", size=9)
    CTR   = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin  = Side(style="thin")
    BRD   = Border(left=thin, right=thin, top=thin, bottom=thin)
    return HDR, Q1F, Q2F, Q3F, Q4F, TOTF, LABF, MODF, FW, FB, CTR, BRD

def _sc(ws, coord, val=None, fill=None, font=None, align=None, border=None):
    c = ws[coord]
    if val is not None: c.value = val
    if fill:   c.fill   = fill
    if font:   c.font   = font
    if align:  c.alignment = align
    if border: c.border = border
    return c

# ── Pages ─────────────────────────────────────────────────────────────────────

@marks_bp.route("/", methods=["GET"])
@login_required
def entry_page():
    if not _fac_ok(): return jsonify({"error": "Forbidden"}), 403
    return render_template("marks/entry.html")

@marks_bp.route("/direct", methods=["GET"])
@login_required
def direct_page():
    if not _fac_ok(): return jsonify({"error": "Forbidden"}), 403
    return render_template("marks/direct.html")

@marks_bp.route("/my", methods=["GET"])
@login_required
def my_marks_page():
    if not current_user.is_student: return jsonify({"error": "Forbidden"}), 403
    return render_template("marks/my_marks.html")

@marks_bp.route("/api/my", methods=["GET"])
@login_required
def api_my_marks():
    if not current_user.is_student: return jsonify({"success": False}), 403
    recs = (InternalMarks.query
            .filter_by(student_id=current_user.id)
            .order_by(InternalMarks.semester, InternalMarks.subject).all())
    return jsonify({"success": True, "marks": [m.to_dict() for m in recs], "total": len(recs)})

# ── Direct entry: load students with existing marks ───────────────────────────

@marks_bp.route("/api/students", methods=["GET"])
@login_required
def api_students():
    """Return students for sem+section with any existing marks for this subject."""
    if not _fac_ok(): return jsonify({"success": False}), 403
    semester     = request.args.get("semester", type=int)
    section      = request.args.get("section", "").strip()
    subject      = request.args.get("subject", "").strip()
    subject_type = request.args.get("subject_type", "cc_theory").strip()

    if not (semester and section):
        return jsonify({"success": False, "message": "semester and section required"}), 400

    students = (User.query
                .filter_by(role="student", semester=semester, section=section)
                .order_by(User.roll_number).all())

    existing = {}
    if subject:
        existing = {m.student_id: m for m in
                    InternalMarks.query.filter_by(
                        subject=subject, semester=semester, section=section).all()}

    rows = []
    for st in students:
        m = existing.get(st.id)
        row = {
            "student_id":  st.id,
            "roll_number": st.roll_number or st.username or "",
            "name":        st.name,
            "subject_type": subject_type,
        }
        if m:
            row.update(m.to_dict())
        rows.append(row)

    return jsonify({"success": True, "students": rows, "total": len(rows)})


# ── Subjects for a semester (from attendance + existing marks) ────────────────

@marks_bp.route("/api/subjects-by-sem", methods=["GET"])
@login_required
def api_subjects_by_sem():
    """Return distinct subjects for a semester from attendance records + existing marks."""
    if not _fac_ok(): return jsonify({"success": False}), 403
    semester = request.args.get("semester", type=int)
    section  = request.args.get("section", "").strip()
    if not semester:
        return jsonify({"success": False, "message": "semester required"}), 400

    from models.attendance import AttendanceSummary
    import re

    _LAB_SUFFIXES = re.compile(
        r'\s*[\(\[]?\s*(lab(oratory)?|practical[s]?|[\(]?p[\)]?)\s*[\)\]]?\s*$',
        re.IGNORECASE
    )

    def _canonical(name):
        """Return (base_name, is_lab) — strips lab suffix from name."""
        stripped = _LAB_SUFFIXES.sub('', (name or '').strip())
        is_lab   = stripped.lower() != (name or '').strip().lower()
        return stripped.strip() or (name or '').strip(), is_lab

    # Collect all raw attendance subjects first
    att_q = (AttendanceSummary.query
             .with_entities(AttendanceSummary.subject)
             .filter_by(semester=semester))
    if section:
        att_q = att_q.filter_by(section=section)
    raw_att = [row.subject for row in att_q.distinct().all() if row.subject]

    # Determine which canonical names have a lab variant
    lab_bases = set()
    for s in raw_att:
        base, is_lab = _canonical(s)
        if is_lab:
            lab_bases.add(base.lower())

    seen      = set()   # canonical lowercase keys already added
    subjects  = []

    # 1) From AttendanceSummary — use canonical name, skip pure-lab duplicates
    for s in raw_att:
        base, is_lab = _canonical(s)
        key = base.lower()
        if key in seen:
            continue          # already added (e.g. already saw theory entry)
        seen.add(key)

        # Determine subject_type:
        #  - if this IS the lab variant → ipcc_lab
        #  - if this is the theory entry but a lab variant also exists → ipcc_lab
        #  - otherwise look up from existing InternalMarks, default cc_theory
        has_lab = (is_lab or key in lab_bases)

        im = (InternalMarks.query
              .filter_by(semester=semester)
              .filter(InternalMarks.subject.ilike(f'{base}%'))
              .first())

        if im:
            stype  = im.subject_type or ('ipcc_lab' if has_lab else 'cc_theory')
            scode  = im.subject_code or ''
            syear  = im.academic_year or ''
        else:
            stype  = 'ipcc_lab' if has_lab else 'cc_theory'
            scode  = ''
            syear  = ''

        subjects.append({
            'subject':      base,        # use canonical (lab-suffix-stripped) name
            'subject_code': scode,
            'subject_type': stype,
            'academic_year': syear,
        })

    # 2) Merge subjects already in InternalMarks (not yet in attendance)
    im_q = (InternalMarks.query
            .with_entities(
                InternalMarks.subject,
                InternalMarks.subject_code,
                InternalMarks.subject_type,
                InternalMarks.academic_year,
            )
            .filter_by(semester=semester))
    if section:
        im_q = im_q.filter_by(section=section)
    for r in im_q.distinct().all():
        base, _ = _canonical(r.subject or '')
        key = base.lower()
        if key and key not in seen:
            seen.add(key)
            subjects.append({
                'subject':      base,
                'subject_code': r.subject_code or '',
                'subject_type': r.subject_type or 'cc_theory',
                'academic_year': r.academic_year or '',
            })

    subjects.sort(key=lambda x: x['subject'].lower())
    return jsonify({'success': True, 'subjects': subjects})


# ── Direct entry: save marks from JSON (no Excel) ─────────────────────────────

@marks_bp.route("/api/save", methods=["POST"])
@login_required
def api_save():
    """
    Save marks directly from JSON payload — the in-app direct entry mode.
    Body: { subject, subject_code, semester, section, academic_year, subject_type, rows: [...] }
    Each row: { student_id, <mark fields> }
    """
    if not _fac_ok(): return jsonify({"success": False, "message": "Forbidden"}), 403

    data          = request.get_json(force=True) or {}
    subject       = data.get("subject", "").strip()
    subject_code  = data.get("subject_code", "").strip()
    semester      = data.get("semester")
    section       = data.get("section", "").strip()
    academic_year = data.get("academic_year", "").strip()
    subject_type  = data.get("subject_type", "cc_theory").strip()
    rows          = data.get("rows", [])

    if not subject or not semester:
        return jsonify({"success": False, "message": "subject and semester are required"}), 400

    asgn_max = ASSIGN_MAX.get(subject_type, 20)
    saved, errors = 0, []

    for row in rows:
        student_id = row.get("student_id")
        if not student_id:
            continue
        try:
            rec = (InternalMarks.query
                   .filter_by(student_id=student_id, subject=subject, semester=semester)
                   .first() or
                   InternalMarks(student_id=student_id, subject=subject, semester=semester))

            rec.faculty_id    = current_user.id
            rec.subject_code  = subject_code
            rec.section       = section
            rec.academic_year = academic_year
            rec.subject_type  = subject_type
            rec.max_assignment = asgn_max

            if subject_type in ('ipcc_theory', 'cc_theory', 'ipcc_lab'):
                for ia in ('ia1', 'ia2'):
                    for q in ('q1', 'q2', 'q3', 'q4'):
                        for sub in ('a', 'b', 'c', 'd'):
                            attr = f"{ia}_{q}_{sub}"
                            setattr(rec, attr, _f(row.get(attr)))
                rec.assignment = _f(row.get("assignment"))
                if subject_type == 'ipcc_lab':
                    rec.lab_ia = _f(row.get("lab_ia"))
                rec.recompute_totals()

            elif subject_type == 'cc_activity':
                for i in range(1, 6):
                    setattr(rec, f"mod{i}", _f(row.get(f"mod{i}")))

            elif subject_type == 'cc_oe':
                rec.oe_gen = _f(row.get("oe_gen"))
                rec.oe_cie = _f(row.get("oe_cie"))

            db.session.add(rec)
            saved += 1
        except Exception as e:
            errors.append(f"Student {student_id}: {e}")

    try:
        db.session.commit()
    except Exception as e:
        db.session.rollback()
        return jsonify({"success": False, "message": f"DB error: {e}"}), 500

    return jsonify({
        "success": True, "saved": saved, "errors": errors,
        "message": f"Saved marks for {saved} student(s)." +
                   (f" {len(errors)} error(s)." if errors else "")
    })


# ── Download template ─────────────────────────────────────────────────────────

@marks_bp.route("/api/template", methods=["GET"])
@login_required
def api_template():
    """Download Excel template pre-filled with students."""
    if not _fac_ok(): return jsonify({"error": "Forbidden"}), 403

    subject      = request.args.get("subject", "").strip()
    code         = request.args.get("subject_code", "").strip()
    semester     = request.args.get("semester", type=int)
    section      = request.args.get("section", "").strip()
    ay           = request.args.get("academic_year", "").strip()
    subject_type = request.args.get("subject_type", "cc_theory").strip()
    faculty_name = request.args.get("faculty_name", "").strip()

    if not (semester and section):
        return jsonify({"error": "Please select a Semester and Section before downloading the template."}), 400

    students = (User.query
                .filter_by(role="student", semester=semester, section=section)
                .order_by(User.roll_number).all())

    if not students:
        return jsonify({"error": f"No students found for Semester {semester}, Section {section}. "
                                  "Make sure students are registered in the system."}), 404

    existing = {}
    if subject:
        existing = {m.student_id: m for m in
                    InternalMarks.query.filter_by(
                        subject=subject, semester=semester, section=section).all()}

    wb = _build_excel(students, existing, subject, code, semester, section, ay,
                      subject_type, faculty_name)

    out = io.BytesIO()
    wb.save(out); out.seek(0)
    stype_short = subject_type.replace("_", "-")
    fname = f"CIE_{stype_short}_Sem{semester}_{section}_{datetime.now().strftime('%Y%m%d')}.xlsx"
    response = send_file(
        out,
        as_attachment=True,
        download_name=fname,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    # Expose Content-Disposition so the JS fetch can read the filename
    response.headers["Access-Control-Expose-Headers"] = "Content-Disposition"
    return response


# ── Excel builders ────────────────────────────────────────────────────────────

def _build_excel(students, existing, subject, code, semester, section, ay,
                 subject_type, faculty_name=""):
    if subject_type in ('ipcc_theory', 'cc_theory', 'ipcc_lab'):
        return _build_q1q4_excel(students, existing, subject, code, semester,
                                  section, ay, subject_type, faculty_name)
    elif subject_type == 'cc_activity':
        return _build_activity_excel(students, existing, subject, code,
                                      semester, section, ay, faculty_name)
    elif subject_type == 'cc_oe':
        return _build_oe_excel(students, existing, subject, code,
                                semester, section, ay, faculty_name)
    return _build_q1q4_excel(students, existing, subject, code, semester,
                              section, ay, 'cc_theory', faculty_name)


def _build_q1q4_excel(students, existing, subject, code, semester, section, ay,
                       subject_type, faculty_name):
    HDR, Q1F, Q2F, Q3F, Q4F, TOTF, LABF, MODF, FW, FB, CTR, BRD = _mk_styles()
    scale   = SCALE_MAX.get(subject_type, 15)
    asgn    = ASSIGN_MAX.get(subject_type, 20)
    has_lab = (subject_type == 'ipcc_lab')
    last_col = "BA" if has_lab else "AZ"

    wb = Workbook(); ws = wb.active; ws.title = "CIE_Marks"
    ws.merge_cells(f"A1:{last_col}1")
    _sc(ws, "A1", val=f"Faculty Name : {faculty_name}\n\nSubject Name : {subject}  ({code})",
        fill=HDR, font=FW, align=CTR)
    ws.row_dimensions[1].height = 45

    ws.merge_cells("A2:B2"); _sc(ws, "A2", val=f"III_Sem_IA_MARKS {ay}", fill=HDR, font=FW, align=CTR)
    # Merge only the sub-question INPUT columns (C-R = 4 questions × 4 sub-parts)
    # Leaves S-Z free for the individual total-column headers
    ws.merge_cells("C2:R2"); _sc(ws, "C2", val="1st INTERNAL ASSESMENT MARKS", fill=HDR, font=FW, align=CTR)
    for col, lbl in [("S2","Total Marks\nQ.No.1"),("T2","Total Marks Q.No.2"),
                      ("U2","Total Marks Q.No.3"),("V2","Total Marks Q.No.4"),
                      ("W2","Highest Marks\nfrom Module -1"),("X2","Highest Marks\nfrom Module -2"),
                      ("Y2","Total Marks\nOut of 50"),("Z2",f"Scaled Down\nto {scale} M")]:
        _sc(ws, col, val=lbl, fill=TOTF, font=FB, align=CTR)
    # Merge only the sub-question INPUT columns for IA2 (AB-AQ = 4 questions × 4 sub-parts)
    ws.merge_cells("AB2:AQ2"); _sc(ws, "AB2", val="2nd INTERNAL ASSESMENT MARKS", fill=HDR, font=FW, align=CTR)
    for col, lbl in [("AR2","Total Marks\nQ.No.1"),("AS2","Total Marks Q.No.2"),
                      ("AT2","Total Marks Q.No.3"),("AU2","Total Marks Q.No.4"),
                      ("AV2","Highest Marks\nfrom Module -1"),("AW2","Highest Marks\nfrom Module -2"),
                      ("AX2","Total Marks\nOut of 50"),("AY2",f"Scaled Down\nto {scale} M")]:
        _sc(ws, col, val=lbl, fill=TOTF, font=FB, align=CTR)
    _sc(ws, "AZ2", val=f"Assignment\n({asgn})", fill=TOTF, font=FB, align=CTR)
    if has_lab: _sc(ws, "BA2", val="LAB IA\nMarks (25)", fill=LABF, font=FB, align=CTR)
    ws.row_dimensions[2].height = 32

    q_groups_ia1 = [("C3","F3","Question No. 1"),("G3","J3","Question No. 2"),
                    ("K3","N3","Question No. 3"),("O3","R3","Question No. 4")]
    q_groups_ia2 = [("AB3","AE3","Question No. 1"),("AF3","AI3","Question No. 2"),
                    ("AJ3","AM3","Question No. 3"),("AN3","AQ3","Question No. 4")]
    q_fills = [Q1F, Q2F, Q3F, Q4F]
    for groups in (q_groups_ia1, q_groups_ia2):
        for (s,e,lbl), fill in zip(groups, q_fills):
            ws.merge_cells(f"{s}:{e}"); _sc(ws, s, val=lbl, fill=fill, font=FB, align=CTR)
    for col in ["A3","B3","S3","T3","U3","V3","W3","X3","Y3","Z3",
                "AA3","AR3","AS3","AT3","AU3","AV3","AW3","AX3","AY3","AZ3"]:
        _sc(ws, col, val="", fill=HDR, font=FW, align=CTR, border=BRD)
    if has_lab: _sc(ws, "BA3", val="", fill=LABF, font=FB, align=CTR, border=BRD)
    ws.row_dimensions[3].height = 18

    ia1_sub = {"C":"a","D":"b","E":"c","F":"d","G":"a","H":"b","I":"c","J":"d",
               "K":"a","L":"b","M":"c","N":"d","O":"a","P":"b","Q":"c","R":"d"}
    ia2_sub = {"AB":"a","AC":"b","AD":"c","AE":"d","AF":"a","AG":"b","AH":"c","AI":"d",
               "AJ":"a","AK":"b","AL":"c","AM":"d","AN":"a","AO":"b","AP":"c","AQ":"d"}
    q_col_fill = {}
    for col in ["C","D","E","F"]:             q_col_fill[col] = Q1F
    for col in ["G","H","I","J"]:             q_col_fill[col] = Q2F
    for col in ["K","L","M","N"]:             q_col_fill[col] = Q3F
    for col in ["O","P","Q","R"]:             q_col_fill[col] = Q4F
    for col in ["AB","AC","AD","AE"]:         q_col_fill[col] = Q1F
    for col in ["AF","AG","AH","AI"]:         q_col_fill[col] = Q2F
    for col in ["AJ","AK","AL","AM"]:         q_col_fill[col] = Q3F
    for col in ["AN","AO","AP","AQ"]:         q_col_fill[col] = Q4F
    for col, lbl in {**ia1_sub, **ia2_sub}.items():
        _sc(ws, f"{col}4", val=lbl, fill=q_col_fill.get(col, HDR), font=FB, align=CTR, border=BRD)
    for col in ["A4","B4","S4","T4","U4","V4","W4","X4","Y4","Z4",
                "AA4","AR4","AS4","AT4","AU4","AV4","AW4","AX4","AY4","AZ4"]:
        _sc(ws, col, val="", fill=HDR, font=FW, align=CTR, border=BRD)
    if has_lab: _sc(ws, "BA4", val="", fill=LABF, font=FB, align=CTR, border=BRD)
    ws.row_dimensions[4].height = 14

    ws["A5"] = "DATE: "; ws["A6"] = "Course Outcomes"
    ws["A7"] = "Marks of Each Sub Question"
    for col in ["S7","T7","U7","V7","AR7","AS7","AT7","AU7"]: ws[col] = 25.0
    ws.row_dimensions[5].height = ws.row_dimensions[6].height = ws.row_dimensions[7].height = 14

    _sc(ws, "A8", val="USN",          fill=HDR, font=FW, align=CTR, border=BRD)
    _sc(ws, "B8", val="Student Name", fill=HDR, font=FW, align=CTR, border=BRD)
    ws.row_dimensions[8].height = 16

    ws.column_dimensions["A"].width = 14; ws.column_dimensions["B"].width = 24
    ws.column_dimensions["AA"].width = 3
    for col in list("CDEFGHIJKLMNOPQR") + ["AB","AC","AD","AE","AF","AG","AH","AI",
                                            "AJ","AK","AL","AM","AN","AO","AP","AQ"]:
        ws.column_dimensions[col].width = 5
    for col in ["S","T","U","V","W","X","Y","Z","AR","AS","AT","AU","AV","AW","AX","AY","AZ"]:
        ws.column_dimensions[col].width = 9
    if has_lab: ws.column_dimensions["BA"].width = 11

    ia1_cols = {"C":"ia1_q1_a","D":"ia1_q1_b","E":"ia1_q1_c","F":"ia1_q1_d",
                "G":"ia1_q2_a","H":"ia1_q2_b","I":"ia1_q2_c","J":"ia1_q2_d",
                "K":"ia1_q3_a","L":"ia1_q3_b","M":"ia1_q3_c","N":"ia1_q3_d",
                "O":"ia1_q4_a","P":"ia1_q4_b","Q":"ia1_q4_c","R":"ia1_q4_d"}
    ia2_cols = {"AB":"ia2_q1_a","AC":"ia2_q1_b","AD":"ia2_q1_c","AE":"ia2_q1_d",
                "AF":"ia2_q2_a","AG":"ia2_q2_b","AH":"ia2_q2_c","AI":"ia2_q2_d",
                "AJ":"ia2_q3_a","AK":"ia2_q3_b","AL":"ia2_q3_c","AM":"ia2_q3_d",
                "AN":"ia2_q4_a","AO":"ia2_q4_b","AP":"ia2_q4_c","AQ":"ia2_q4_d"}
    col_fill_map = {
        "C":Q1F,"D":Q1F,"E":Q1F,"F":Q1F,"G":Q2F,"H":Q2F,"I":Q2F,"J":Q2F,
        "K":Q3F,"L":Q3F,"M":Q3F,"N":Q3F,"O":Q4F,"P":Q4F,"Q":Q4F,"R":Q4F,
        "AB":Q1F,"AC":Q1F,"AD":Q1F,"AE":Q1F,"AF":Q2F,"AG":Q2F,"AH":Q2F,"AI":Q2F,
        "AJ":Q3F,"AK":Q3F,"AL":Q3F,"AM":Q3F,"AN":Q4F,"AO":Q4F,"AP":Q4F,"AQ":Q4F}

    for r, st in enumerate(students, start=9):
        m = existing.get(st.id)
        ws[f"A{r}"] = st.roll_number or st.username or ""; ws[f"B{r}"] = st.name
        _sc(ws, f"A{r}", border=BRD, align=CTR); _sc(ws, f"B{r}", border=BRD, align=CTR)
        for col, attr in ia1_cols.items():
            _sc(ws, f"{col}{r}", val=getattr(m, attr, None) if m else None, fill=col_fill_map[col], align=CTR, border=BRD)
        for col, attr in ia2_cols.items():
            _sc(ws, f"{col}{r}", val=getattr(m, attr, None) if m else None, fill=col_fill_map[col], align=CTR, border=BRD)
        for f_col, formula in [("S",f"=SUM(C{r}:F{r})"),("T",f"=SUM(G{r}:J{r})"),
                                 ("U",f"=SUM(K{r}:N{r})"),("V",f"=SUM(O{r}:R{r})"),
                                 ("W",f"=MAX(S{r},T{r})"),("X",f"=MAX(U{r},V{r})"),
                                 ("Y",f"=W{r}+X{r}"),("Z",f"=ROUND(Y{r}*{scale}/50,1)")]:
            ws[f"{f_col}{r}"] = formula
            _sc(ws, f"{f_col}{r}", fill=TOTF, font=Font(bold=True, size=9), align=CTR, border=BRD)
        for f_col, formula in [("AR",f"=SUM(AB{r}:AE{r})"),("AS",f"=SUM(AF{r}:AI{r})"),
                                 ("AT",f"=SUM(AJ{r}:AM{r})"),("AU",f"=SUM(AN{r}:AQ{r})"),
                                 ("AV",f"=MAX(AR{r},AS{r})"),("AW",f"=MAX(AT{r},AU{r})"),
                                 ("AX",f"=AV{r}+AW{r}"),("AY",f"=ROUND(AX{r}*{scale}/50,1)")]:
            ws[f"{f_col}{r}"] = formula
            _sc(ws, f"{f_col}{r}", fill=TOTF, font=Font(bold=True, size=9), align=CTR, border=BRD)
        ws[f"AZ{r}"] = m.assignment if m else None
        _sc(ws, f"AZ{r}", fill=TOTF, font=Font(bold=True, size=9), align=CTR, border=BRD)
        if has_lab:
            ws[f"BA{r}"] = m.lab_ia if m else None
            _sc(ws, f"BA{r}", fill=LABF, font=Font(bold=True, size=9), align=CTR, border=BRD)
        ws.row_dimensions[r].height = 16

    ws.freeze_panes = "C9"
    return wb


def _build_activity_excel(students, existing, subject, code, semester, section, ay, faculty_name):
    HDR, Q1F, Q2F, Q3F, Q4F, TOTF, LABF, MODF, FW, FB, CTR, BRD = _mk_styles()
    wb = Workbook(); ws = wb.active; ws.title = "CIE_Activity"
    ws.merge_cells("A1:H1")
    _sc(ws, "A1", val=f"Faculty Name : {faculty_name}\n\nSubject Name : {subject}  ({code})", fill=HDR, font=FW, align=CTR)
    ws.row_dimensions[1].height = 45
    ws.merge_cells("A2:B2"); _sc(ws, "A2", val=f"III_Sem_IA_MARKS {ay}", fill=HDR, font=FW, align=CTR)
    for i, col in enumerate(["C","D","E","F","G"], 1):
        _sc(ws, f"{col}2", val=f"MODULE {i}=20M", fill=MODF[i-1], font=FB, align=CTR, border=BRD)
    _sc(ws, "H2", val="TOTAL-100", fill=TOTF, font=FB, align=CTR, border=BRD)
    ws.row_dimensions[2].height = 22
    for row in range(3, 8): ws.row_dimensions[row].height = 14
    ws["A5"] = "DATE: "; ws["A6"] = "Course Outcomes"; ws["A7"] = "Marks of Each Sub Question"
    _sc(ws, "A8", val="USN", fill=HDR, font=FW, align=CTR, border=BRD)
    _sc(ws, "B8", val="Student Name", fill=HDR, font=FW, align=CTR, border=BRD)
    ws.row_dimensions[8].height = 16
    ws.column_dimensions["A"].width = 14; ws.column_dimensions["B"].width = 26
    for col in ["C","D","E","F","G"]: ws.column_dimensions[col].width = 12
    ws.column_dimensions["H"].width = 14
    for r, st in enumerate(students, start=9):
        m = existing.get(st.id)
        ws[f"A{r}"] = st.roll_number or st.username or ""; ws[f"B{r}"] = st.name
        _sc(ws, f"A{r}", border=BRD, align=CTR); _sc(ws, f"B{r}", border=BRD, align=CTR)
        for i, col in enumerate(["C","D","E","F","G"], 1):
            _sc(ws, f"{col}{r}", val=getattr(m, f"mod{i}", None) if m else None, fill=MODF[i-1], align=CTR, border=BRD)
        ws[f"H{r}"] = f"=SUM(C{r}:G{r})"
        _sc(ws, f"H{r}", fill=TOTF, font=Font(bold=True, size=9), align=CTR, border=BRD)
        ws.row_dimensions[r].height = 16
    ws.freeze_panes = "C9"
    return wb


def _build_oe_excel(students, existing, subject, code, semester, section, ay, faculty_name):
    HDR, Q1F, Q2F, Q3F, Q4F, TOTF, LABF, MODF, FW, FB, CTR, BRD = _mk_styles()
    wb = Workbook(); ws = wb.active; ws.title = "CIE_OpenElective"
    ws.merge_cells("A1:E1")
    _sc(ws, "A1", val=f"Faculty Name : {faculty_name}\n\nSubject Name : {subject}  ({code})", fill=HDR, font=FW, align=CTR)
    ws.row_dimensions[1].height = 45
    ws.merge_cells("A2:B2"); _sc(ws, "A2", val=f"III_Sem_IA_MARKS {ay}", fill=HDR, font=FW, align=CTR)
    _sc(ws, "C2", val="Gen\n 30 M", fill=Q1F, font=FB, align=CTR, border=BRD)
    _sc(ws, "D2", val="CIE 20",    fill=Q2F, font=FB, align=CTR, border=BRD)
    _sc(ws, "E2", val="Total 50",  fill=TOTF, font=FB, align=CTR, border=BRD)
    ws.row_dimensions[2].height = 22
    for row in range(3, 8): ws.row_dimensions[row].height = 14
    ws["A5"] = "DATE: "; ws["A6"] = "Course Outcomes"; ws["A7"] = "Marks of Each Sub Question"
    _sc(ws, "A8", val="USN", fill=HDR, font=FW, align=CTR, border=BRD)
    _sc(ws, "B8", val="Student Name", fill=HDR, font=FW, align=CTR, border=BRD)
    ws.row_dimensions[8].height = 16
    ws.column_dimensions["A"].width = 14; ws.column_dimensions["B"].width = 26
    ws.column_dimensions["C"].width = 14; ws.column_dimensions["D"].width = 12; ws.column_dimensions["E"].width = 12
    for r, st in enumerate(students, start=9):
        m = existing.get(st.id)
        ws[f"A{r}"] = st.roll_number or st.username or ""; ws[f"B{r}"] = st.name
        _sc(ws, f"A{r}", border=BRD, align=CTR); _sc(ws, f"B{r}", border=BRD, align=CTR)
        ws[f"C{r}"] = m.oe_gen if m else None; ws[f"D{r}"] = m.oe_cie if m else None
        _sc(ws, f"C{r}", fill=Q1F, align=CTR, border=BRD); _sc(ws, f"D{r}", fill=Q2F, align=CTR, border=BRD)
        ws[f"E{r}"] = f"=C{r}+D{r}"
        _sc(ws, f"E{r}", fill=TOTF, font=Font(bold=True, size=9), align=CTR, border=BRD)
        ws.row_dimensions[r].height = 16
    ws.freeze_panes = "C9"
    return wb


# ── Upload filled Excel ───────────────────────────────────────────────────────

@marks_bp.route("/api/upload", methods=["POST"])
@login_required
def api_upload():
    if not _fac_ok(): return jsonify({"success": False, "message": "Forbidden"}), 403
    if "file" not in request.files:
        return jsonify({"success": False, "message": "No file uploaded"}), 400

    subject       = request.form.get("subject", "").strip()
    subject_code  = request.form.get("subject_code", "").strip()
    semester      = request.form.get("semester", type=int)
    section       = request.form.get("section", "").strip()
    academic_year = request.form.get("academic_year", "").strip()
    subject_type  = request.form.get("subject_type", "cc_theory").strip()

    file = request.files["file"]
    try:
        wb2 = load_workbook(file.stream, data_only=True)
        ws2 = wb2.active
    except Exception as e:
        return jsonify({"success": False, "message": f"Cannot read file: {e}"}), 400

    if subject_type in ('ipcc_theory', 'cc_theory', 'ipcc_lab'):
        saved, errors = _parse_q1q4(ws2, subject, subject_code, semester, section, academic_year, subject_type)
    elif subject_type == 'cc_activity':
        saved, errors = _parse_activity(ws2, subject, subject_code, semester, section, academic_year)
    elif subject_type == 'cc_oe':
        saved, errors = _parse_oe(ws2, subject, subject_code, semester, section, academic_year)
    else:
        return jsonify({"success": False, "message": f"Unknown subject_type: {subject_type}"}), 400

    try:
        db.session.commit()
    except Exception as e:
        db.session.rollback()
        return jsonify({"success": False, "message": f"DB error: {e}"}), 500

    return jsonify({
        "success": True, "saved": saved, "errors": errors,
        "message": f"Saved marks for {saved} student(s)." +
                   (f" {len(errors)} error(s)." if errors else "")
    })


def _get_or_create_record(usn, roll, subject, semester):
    st = None
    if usn:
        st = User.query.filter_by(username=str(usn).strip(), role="student").first()
    if not st and roll:
        st = User.query.filter_by(roll_number=str(roll).strip(), role="student").first()
    if not st:
        return None, None
    rec = (InternalMarks.query
           .filter_by(student_id=st.id, subject=subject, semester=semester).first()
           or InternalMarks(student_id=st.id, subject=subject, semester=semester))
    return st, rec


def _parse_q1q4(ws, subject, subject_code, semester, section, academic_year, subject_type):
    has_lab  = (subject_type == 'ipcc_lab')
    asgn_max = ASSIGN_MAX.get(subject_type, 20)
    ia1_map  = {"C":"ia1_q1_a","D":"ia1_q1_b","E":"ia1_q1_c","F":"ia1_q1_d",
                "G":"ia1_q2_a","H":"ia1_q2_b","I":"ia1_q2_c","J":"ia1_q2_d",
                "K":"ia1_q3_a","L":"ia1_q3_b","M":"ia1_q3_c","N":"ia1_q3_d",
                "O":"ia1_q4_a","P":"ia1_q4_b","Q":"ia1_q4_c","R":"ia1_q4_d"}
    ia2_map  = {"AB":"ia2_q1_a","AC":"ia2_q1_b","AD":"ia2_q1_c","AE":"ia2_q1_d",
                "AF":"ia2_q2_a","AG":"ia2_q2_b","AH":"ia2_q2_c","AI":"ia2_q2_d",
                "AJ":"ia2_q3_a","AK":"ia2_q3_b","AL":"ia2_q3_c","AM":"ia2_q3_d",
                "AN":"ia2_q4_a","AO":"ia2_q4_b","AP":"ia2_q4_c","AQ":"ia2_q4_d"}
    saved, errors = 0, []
    for row in range(9, ws.max_row + 1):
        usn = ws[f"A{row}"].value; name = ws[f"B{row}"].value
        if not usn and not name: continue
        st, rec = _get_or_create_record(usn, usn, subject, semester)
        if not st: errors.append(f"Row {row}: USN '{usn}' not found"); continue
        try:
            rec.faculty_id = current_user.id; rec.subject_code = subject_code
            rec.section = section; rec.academic_year = academic_year
            rec.subject_type = subject_type; rec.max_assignment = asgn_max
            for col, attr in ia1_map.items(): setattr(rec, attr, _f(ws[f"{col}{row}"].value))
            for col, attr in ia2_map.items(): setattr(rec, attr, _f(ws[f"{col}{row}"].value))
            rec.assignment = _f(ws[f"AZ{row}"].value)
            if has_lab: rec.lab_ia = _f(ws[f"BA{row}"].value)
            rec.recompute_totals(); db.session.add(rec); saved += 1
        except Exception as e: errors.append(f"Row {row}: {e}")
    return saved, errors


def _parse_activity(ws, subject, subject_code, semester, section, academic_year):
    saved, errors = 0, []
    for row in range(9, ws.max_row + 1):
        usn = ws[f"A{row}"].value; name = ws[f"B{row}"].value
        if not usn and not name: continue
        st, rec = _get_or_create_record(usn, usn, subject, semester)
        if not st: errors.append(f"Row {row}: USN '{usn}' not found"); continue
        try:
            rec.faculty_id = current_user.id; rec.subject_code = subject_code
            rec.section = section; rec.academic_year = academic_year; rec.subject_type = 'cc_activity'
            for i, col in enumerate(["C","D","E","F","G"], 1):
                setattr(rec, f"mod{i}", _f(ws[f"{col}{row}"].value))
            db.session.add(rec); saved += 1
        except Exception as e: errors.append(f"Row {row}: {e}")
    return saved, errors


def _parse_oe(ws, subject, subject_code, semester, section, academic_year):
    saved, errors = 0, []
    for row in range(9, ws.max_row + 1):
        usn = ws[f"A{row}"].value; name = ws[f"B{row}"].value
        if not usn and not name: continue
        st, rec = _get_or_create_record(usn, usn, subject, semester)
        if not st: errors.append(f"Row {row}: USN '{usn}' not found"); continue
        try:
            rec.faculty_id = current_user.id; rec.subject_code = subject_code
            rec.section = section; rec.academic_year = academic_year; rec.subject_type = 'cc_oe'
            rec.oe_gen = _f(ws[f"C{row}"].value); rec.oe_cie = _f(ws[f"D{row}"].value)
            db.session.add(rec); saved += 1
        except Exception as e: errors.append(f"Row {row}: {e}")
    return saved, errors


# ── View roster ───────────────────────────────────────────────────────────────

@marks_bp.route("/api/roster", methods=["GET"])
@login_required
def api_roster():
    if not _fac_ok(): return jsonify({"success": False}), 403
    subject  = request.args.get("subject", "").strip()
    semester = request.args.get("semester", type=int)
    section  = request.args.get("section", "").strip()
    if not (subject and semester and section):
        return jsonify({"success": False, "message": "Missing params"}), 400
    marks = (InternalMarks.query
             .filter_by(subject=subject, semester=semester, section=section)
             .join(User, InternalMarks.student_id == User.id)
             .order_by(User.roll_number).all())
    return jsonify({"success": True, "rows": [m.to_dict() for m in marks],
                    "total": len(marks)})

